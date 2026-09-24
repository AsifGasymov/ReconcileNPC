"""XMGate · Duplicates — merge gateway CSV exports and flag pay ids from a duplicates list."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment

from .excel_styles import (
    ALT, BORDER, DARK_BLUE, ORANGE_FILL, WHITE,
    data_font, header_font, style_cell,
)
from .io import read_table

LogFn = Callable[[str], None]

NUMBER_FMT = "#,##0.00"
DATE_FMT   = "DD.MM.YYYY"

PAY_ID_COL = "pay id"
DUP_COL    = "Duplicate"
# merchant_order_id "payXXXX…" → "pay_" + chars 4..27 (first 27 chars, "_" after "pay")
PAY_ID_LEN = 27

AMOUNT_COLS = {"amount", "accounting_amount"}
WIDE_COLS   = {PAY_ID_COL: 32, "merchant_order_id": 34, "descriptor": 36,
               "remote_id": 34, "order_id": 22, "terminal": 20}


@dataclass
class XMGateResult:
    total_rows: int
    dup_pay_ids: int
    dup_rows: int
    dup_charge_count: int
    dup_charge_amount: float
    dup_list_size: int
    out_path: str


def make_pay_id(merchant_order_id: str) -> str:
    s = str(merchant_order_id or "").strip()
    if not s.lower().startswith("pay"):
        return ""
    return "pay_" + s[3:PAY_ID_LEN]


def run_xmgate_dupes(
    *,
    export_paths: list[str],
    dupes_path: str,
    out_dir: str,
    log: Optional[LogFn] = None,
) -> XMGateResult:
    def _log(msg: str) -> None:
        if log:
            log(msg)

    _log(f"Reading {len(export_paths)} gateway export(s)…")
    parts = [read_table(p).fillna("") for p in export_paths]
    df = pd.concat(parts, ignore_index=True, sort=False).fillna("")
    if "merchant_order_id" not in df.columns:
        raise ValueError("Column 'merchant_order_id' not found in exports")

    df.insert(0, PAY_ID_COL, df["merchant_order_id"].map(make_pay_id))
    if "created" in df.columns:
        df["created"] = pd.to_datetime(df["created"], errors="coerce").dt.normalize()
    for c in AMOUNT_COLS & set(df.columns):
        df[c] = pd.to_numeric(df[c].astype(str).str.replace(",", ".", regex=False),
                              errors="coerce")

    _log("Reading duplicates list…")
    dl = read_table(dupes_path)
    id_col = next((c for c in dl.columns if c.strip().lower() == PAY_ID_COL), dl.columns[0])
    dup_ids = set(dl[id_col].dropna().astype(str).str.strip()) - {""}
    _log(f"Export rows: {len(df)}  |  Duplicate list ids: {len(dup_ids)}")

    _log("Checking duplicates…")
    is_dup = df[PAY_ID_COL].isin(dup_ids)
    df.insert(1, DUP_COL, is_dup.map({True: "Yes", False: "No"}))
    dups = df[is_dup].copy()

    charges = dups[dups.get("type", pd.Series("", index=dups.index)) == "charge"]
    dup_charge_amount = float(charges["amount"].sum()) if "amount" in charges else 0.0
    n_dup_ids = dups[PAY_ID_COL].nunique()
    _log(f"Duplicate pay ids: {n_dup_ids}  |  rows: {len(dups)}  |  "
         f"charges: {len(charges)} / {dup_charge_amount:,.2f}")

    # ── Write Excel ──────────────────────────────────────────────────────────
    _log("Writing output…")
    hfont = header_font()
    dfont = data_font()
    cols = list(df.columns)

    def _write_sheet(ws, data: pd.DataFrame) -> None:
        for ci, name in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=ci, value=name)
            cell.fill = DARK_BLUE
            cell.font = hfont
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[cell.column_letter].width = WIDE_COLS.get(name, 14)
        ws.row_dimensions[1].height = 20
        ws.freeze_panes = "C2"

        for ri, row in enumerate(data.itertuples(index=False), start=2):
            alt = ALT if ri % 2 == 0 else WHITE
            for ci, (name, val) in enumerate(zip(cols, row), start=1):
                if pd.isna(val) or val == "":
                    val = None
                elif isinstance(val, pd.Timestamp):
                    val = val.to_pydatetime()
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.font = dfont
                cell.border = BORDER
                cell.fill = ORANGE_FILL if (name == DUP_COL and val == "Yes") else alt
                if name in AMOUNT_COLS and val is not None:
                    cell.number_format = NUMBER_FMT
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif name == "created" and val is not None:
                    cell.number_format = DATE_FMT
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
        if len(data):
            ws.auto_filter.ref = ws.dimensions

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "All"
    _write_sheet(ws1, df)

    ws2 = wb.create_sheet("Duplicates")
    _write_sheet(ws2, dups)

    ws3 = wb.create_sheet("Summary")
    for ci, (hdr, width) in enumerate([("Metric", 34), ("Value", 18)], start=1):
        cell = ws3.cell(row=1, column=ci, value=hdr)
        cell.fill = DARK_BLUE
        cell.font = hfont
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws3.column_dimensions[cell.column_letter].width = width

    summary_rows: list[tuple[str, object]] = [
        ("Export rows (merged)",          len(df)),
        ("Ids in duplicates list",        len(dup_ids)),
        ("Duplicate pay ids found",       n_dup_ids),
        ("Duplicate rows",                len(dups)),
        ("Duplicate charges",             len(charges)),
        ("Duplicate charge amount",       round(dup_charge_amount, 2)),
    ]
    if len(dups) and "terminal" in dups.columns:
        for term, n in dups.groupby("terminal")[PAY_ID_COL].nunique().sort_values(
                ascending=False).items():
            summary_rows.append((f"  — {term}", int(n)))
    for ri, (label, val) in enumerate(summary_rows, start=2):
        fill = ALT if ri % 2 == 0 else WHITE
        style_cell(ws3, ri, 1, label, fill=fill, font=dfont)
        style_cell(ws3, ri, 2, val, fill=fill, font=dfont, align="right",
                   number_format=NUMBER_FMT if isinstance(val, float) else None)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    _out_dir = Path(out_dir)
    _out_dir.mkdir(parents=True, exist_ok=True)
    out_path = str(_out_dir.resolve() / f"xmgate_duplicates_{ts}.xlsx")
    wb.save(out_path)
    _log(f"Saved → {out_path}")

    return XMGateResult(
        total_rows=len(df),
        dup_pay_ids=n_dup_ids,
        dup_rows=len(dups),
        dup_charge_count=len(charges),
        dup_charge_amount=round(dup_charge_amount, 2),
        dup_list_size=len(dup_ids),
        out_path=out_path,
    )
